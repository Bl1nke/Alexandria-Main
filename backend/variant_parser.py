"""
Parser for EGE variant files.

Expected NeoFamily-like structure:
  - Text 1 before task 1
  - Tasks 1-22
  - Text 2 after task 22 and before task 23
  - Tasks 23-27

The important part is marker detection. A variant contains many numbered
answer options ("1.", "2.", "3."), so we only accept task markers that look
like real task headings and appear in the expected order.
"""

import io
import json
import os
import re
from typing import Dict, List, Optional, Tuple

import fitz
import pdfplumber
from PIL import Image

try:
    from docx import Document as DocxDocument

    DOCX_OK = True
except ImportError:
    DOCX_OK = False


TASK_TYPES = {
    0: "text",
    -1: "text",
    1: "union_selection",
    2: "lexical_meaning",
    3: "text_analysis",
    4: "stress",
    5: "paronym",
    6: "word_correction",
    7: "word_form",
    8: "grammar_error",
    9: "root_vowel",
    10: "prefix",
    11: "suffix",
    12: "verb_ending",
    13: "ne_with_word",
    14: "union_spelling",
    15: "nn_spelling",
    16: "punctuation_one_comma",
    17: "punctuation_participles",
    18: "punctuation_intro",
    19: "punctuation_complex",
    20: "punctuation_complex2",
    21: "dash_usage",
    22: "expressive_means",
    23: "content_analysis",
    24: "text_features",
    25: "phraseology",
    26: "text_connection",
    27: "essay",
}

TASK_NAMES = {
    "text": "Текст",
    "union_selection": "Выбор союза",
    "lexical_meaning": "Лексика",
    "text_analysis": "Анализ текста",
    "stress": "Ударение",
    "paronym": "Паронимы",
    "word_correction": "Исправление слова",
    "word_form": "Форма слова",
    "grammar_error": "Грамматика",
    "root_vowel": "Чередующиеся гласные",
    "prefix": "Приставки",
    "suffix": "Суффиксы",
    "verb_ending": "Окончания глаголов",
    "ne_with_word": "НЕ со словами",
    "union_spelling": "Союзы",
    "nn_spelling": "НН",
    "punctuation_one_comma": "Одна запятая",
    "punctuation_participles": "Обороты",
    "punctuation_intro": "Вводные слова",
    "punctuation_complex": "Сложное предложение",
    "punctuation_complex2": "Сложное с И",
    "dash_usage": "Тире",
    "expressive_means": "Средства выразительности",
    "content_analysis": "Содержание",
    "text_features": "Особенности текста",
    "phraseology": "Фразеологизмы",
    "text_connection": "Средства связи",
    "essay": "Сочинение",
    "unknown": "Задание",
}

TOP_MARGIN = 8
BOTTOM_MARGIN = 4
DPI = 140
DRAWING_SPACE_PX = 180
TASK_NUMBERS = set(range(1, 28))


def parse_answers(path: str) -> Dict[str, str]:
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "docx":
        return _parse_answers_docx(path)
    return _parse_answers_pdf(path)


def _parse_answers_pdf(path: str) -> Dict[str, str]:
    answers: Dict[str, str] = {}
    skip = {"№ п/п", "№ задания", "ответ", "answer", "решение:", "источник:"}

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                for row in table:
                    if not row or len(row) < 2:
                        continue
                    cols = [c.strip() if c else "" for c in row]
                    num_str = cols[0]
                    answer = cols[-1] if len(cols) >= 3 else cols[1]
                    if not num_str or not answer:
                        continue
                    if num_str.lower() in skip or answer.lower() in skip:
                        continue
                    if re.match(r"^\d{1,2}$", num_str):
                        answers[num_str] = answer.strip()

    return answers or _parse_answers_text(path)


def _parse_answers_text(path: str) -> Dict[str, str]:
    answers: Dict[str, str] = {}
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").splitlines():
                line = line.strip()
                m = re.match(r"^(\d{1,2})\s+\d{4,}\s+(.+)$", line)
                if m:
                    answers.setdefault(m.group(1), m.group(2).strip())
                    continue
                m = re.match(r"^(\d{1,2})\s{2,}(.+)$", line)
                if m:
                    ans = m.group(2).strip()
                    if ans.lower() not in ("ответ", "answer"):
                        answers.setdefault(m.group(1), ans)
    return answers


def _parse_answers_docx(path: str) -> Dict[str, str]:
    if not DOCX_OK:
        raise ValueError("python-docx не установлен")

    doc = DocxDocument(path)
    answers: Dict[str, str] = {}
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 2 and re.match(r"^\d{1,2}$", cells[0]):
                answers[cells[0]] = cells[-1]
    return answers


def _crop(
    fitz_doc,
    page,
    y0: float,
    y1: float,
    dpi: int = DPI,
    top_margin: int = TOP_MARGIN,
    bottom_margin: int = BOTTOM_MARGIN,
) -> bytes:
    width = float(page.width)
    height = float(page.height)
    y0 = max(0.0, y0 - top_margin)
    y1 = min(height, y1 + bottom_margin)

    if y1 <= y0:
        y1 = min(height, y0 + 1)

    fitz_page = fitz_doc[page.page_number - 1]
    clip = fitz.Rect(0, y0, width, y1)
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    pix = fitz_page.get_pixmap(matrix=matrix, clip=clip, alpha=False)
    img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _full_page(fitz_doc, page, dpi: int = DPI) -> bytes:
    fitz_page = fitz_doc[page.page_number - 1]
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    pix = fitz_page.get_pixmap(matrix=matrix, alpha=False)
    img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _vstack(parts: List[bytes]) -> bytes:
    if len(parts) == 1:
        return parts[0]

    images = [_trim_vertical_whitespace(Image.open(io.BytesIO(part)).convert("RGB")) for part in parts]
    width = max(img.width for img in images)
    height = sum(img.height for img in images)
    out = Image.new("RGB", (width, height), "white")

    y = 0
    for img in images:
        out.paste(img, (0, y))
        y += img.height

    buf = io.BytesIO()
    out.save(buf, "PNG")
    return buf.getvalue()


def _trim_vertical_whitespace(image: Image.Image, threshold: int = 248, padding: int = 14) -> Image.Image:
    pixels = image.convert("L")
    width, height = pixels.size
    top = 0
    bottom = height - 1

    def row_has_ink(y: int) -> bool:
        for x in range(0, width, 4):
            if pixels.getpixel((x, y)) < threshold:
                return True
        return False

    while top < height and not row_has_ink(top):
        top += 1
    while bottom > top and not row_has_ink(bottom):
        bottom -= 1

    if top >= bottom:
        return image

    top = max(0, top - padding)
    bottom = min(height - 1, bottom + padding)
    return image.crop((0, top, width, bottom + 1))


def _compact_vertical_whitespace(data: bytes, threshold: int = 248, max_blank: int = 70) -> bytes:
    image = Image.open(io.BytesIO(data)).convert("RGB")
    gray = image.convert("L")
    width, height = gray.size
    rows_to_keep: List[int] = []
    blank_run: List[int] = []

    def row_has_ink(y: int) -> bool:
        for x in range(0, width, 4):
            if gray.getpixel((x, y)) < threshold:
                return True
        return False

    def flush_blank_run() -> None:
        if not blank_run:
            return
        if len(blank_run) <= max_blank:
            rows_to_keep.extend(blank_run)
        else:
            keep = max_blank // 2
            rows_to_keep.extend(blank_run[:keep])
            rows_to_keep.extend(blank_run[-(max_blank - keep):])
        blank_run.clear()

    for y in range(height):
        if row_has_ink(y):
            flush_blank_run()
            rows_to_keep.append(y)
        else:
            blank_run.append(y)
    flush_blank_run()

    if len(rows_to_keep) == height:
        return data

    compact = Image.new("RGB", (width, len(rows_to_keep)), "white")
    for new_y, old_y in enumerate(rows_to_keep):
        compact.paste(image.crop((0, old_y, width, old_y + 1)), (0, new_y))

    buf = io.BytesIO()
    compact.save(buf, "PNG")
    return buf.getvalue()


def _add_bottom_drawing_space(data: bytes, height: int = DRAWING_SPACE_PX) -> bytes:
    image = Image.open(io.BytesIO(data)).convert("RGB")
    out = Image.new("RGB", (image.width, image.height + height), "white")
    out.paste(image, (0, 0))

    buf = io.BytesIO()
    out.save(buf, "PNG")
    return buf.getvalue()


def _line_text(words: List[dict]) -> str:
    return " ".join(word["text"] for word in sorted(words, key=lambda w: w["x0"]))


def _group_words_by_line(words: List[dict]) -> List[List[dict]]:
    lines: Dict[int, List[dict]] = {}
    for word in words:
        y = round(float(word["top"]) / 3) * 3
        lines.setdefault(y, []).append(word)
    return [sorted(lines[y], key=lambda w: w["x0"]) for y in sorted(lines)]


def _font_size(word: dict) -> float:
    return float(word.get("bottom", 0)) - float(word.get("top", 0))


def _line_has_bold_font(words: List[dict]) -> bool:
    for word in words[:8]:
        name = str(word.get("fontname", "")).lower()
        if "bold" in name or "black" in name or "heavy" in name:
            return True
    return False


def _looks_like_task_heading(words: List[dict], expected_num: int) -> bool:
    if not words:
        return False

    first = words[0]["text"].strip()
    if first != f"{expected_num}.":
        return False

    line = _line_text(words)
    if re.match(r"^\d+\.\s*$", line):
        return False

    x0 = float(words[0]["x0"])
    size = _font_size(words[0])
    line_upper = line.upper()
    heading_tokens = (
        "САМОСТОЯТЕЛЬНО",
        "В ТЕКСТЕ",
        "УКАЖИТЕ",
        "В ОДНОМ",
        "ОТРЕДАКТИРУЙТЕ",
        "ИСПРАВЬТЕ",
        "ОБРАЗУЙТЕ",
        "УСТАНОВИТЕ",
        "ОПРЕДЕЛИТЕ",
        "ВЫПИШИТЕ",
        "РАССТАВЬТЕ",
        "НАЙДИТЕ",
        "ПРОЧИТАЙТЕ",
        "НАПИШИТЕ",
        "СФОРМУЛИРУЙТЕ",
        "ПРОКОММЕНТИРУЙТЕ",
    )

    # Numbered answer options can also be bold and close to the left margin.
    # Accept only lines that contain wording typical for an actual task header.
    if x0 > 90:
        return False
    return size >= 10 and any(token in line_upper for token in heading_tokens)


def _find_task_markers(pdf) -> List[Tuple[int, int, float]]:
    marks: List[Tuple[int, int, float]] = []
    expected = 1

    for page_index, page in enumerate(pdf.pages):
        words = page.extract_words(extra_attrs=["fontname", "size"])
        for line_words in _group_words_by_line(words):
            if expected > 27:
                return marks
            if _looks_like_task_heading(line_words, expected):
                marks.append((expected, page_index, float(line_words[0]["top"])))
                expected += 1

    return marks


def _find_answer_line_after(pdf, page_start: int, y_start: float, page_end: int, y_end: float) -> Optional[Tuple[int, float]]:
    for page_index in range(page_start, min(page_end, len(pdf.pages) - 1) + 1):
        page = pdf.pages[page_index]
        words = page.extract_words(extra_attrs=["fontname", "size"])
        for line_words in _group_words_by_line(words):
            y = float(line_words[0]["top"])
            if page_index == page_start and y <= y_start:
                continue
            if page_index == page_end and y >= y_end:
                continue
            text = _line_text(line_words).upper().replace(" ", "")
            if text.startswith("ОТВЕТ:") or text.startswith("ОТВЕТ"):
                return page_index, float(line_words[-1]["bottom"])
    return None


def _task_end_before_next_marker(
    pdf,
    start_page: int,
    start_y: float,
    next_page: int,
    next_y: float,
) -> Tuple[int, float]:
    answer_end = _find_answer_line_after(pdf, start_page, start_y, next_page, next_y)
    if answer_end:
        return answer_end
    return next_page, next_y


def _find_numbered_text_start_between(
    pdf,
    page_start: int,
    y_start: float,
    page_end: int,
    y_end: float,
) -> Optional[Tuple[int, float]]:
    for page_index in range(page_start, min(page_end, len(pdf.pages) - 1) + 1):
        page = pdf.pages[page_index]
        words = page.extract_words(extra_attrs=["fontname", "size"])
        for line_words in _group_words_by_line(words):
            y = float(line_words[0]["top"])
            if page_index == page_start and y <= y_start:
                continue
            if page_index == page_end and y >= y_end:
                continue
            text = _line_text(line_words).strip()
            if re.match(r"^\(?1\)", text):
                return page_index, y
    return None


def _cut_range(
    fitz_doc,
    pdf,
    start_page: int,
    start_y: float,
    end_page: int,
    end_y: float,
    end_margin: int = BOTTOM_MARGIN,
) -> bytes:
    if start_page == end_page:
        return _crop(fitz_doc, pdf.pages[start_page], start_y, end_y, bottom_margin=end_margin)

    parts = [_crop(fitz_doc, pdf.pages[start_page], start_y, float(pdf.pages[start_page].height))]
    for page_index in range(start_page + 1, end_page):
        parts.append(_full_page(fitz_doc, pdf.pages[page_index]))
    parts.append(_crop(fitz_doc, pdf.pages[end_page], 0.0, end_y, top_margin=0, bottom_margin=end_margin))
    return _vstack(parts)


def extract_task_images_pdf(path: str) -> Dict[int, bytes]:
    with pdfplumber.open(path) as pdf, fitz.open(path) as fitz_doc:
        marks = _find_task_markers(pdf)
        if not marks:
            return {i: _full_page(fitz_doc, page) for i, page in enumerate(pdf.pages, 1)}

        if len(marks) < 27:
            found = ", ".join(str(num) for num, _, _ in marks)
            raise ValueError(f"Найдены не все задания. Найдено: {found}")

        by_num = {num: (page_index, y) for num, page_index, y in marks}
        task_images: Dict[int, bytes] = {}

        task1_page, task1_y = by_num[1]
        text1_image = _cut_range(fitz_doc, pdf, 0, 0.0, task1_page, task1_y, end_margin=0)

        task22_page, task22_y = by_num[22]
        task23_page, task23_y = by_num[23]
        text2_start = _find_numbered_text_start_between(pdf, task22_page, task22_y, task23_page, task23_y)
        if not text2_start:
            text2_start = (task23_page, 0.0) if task23_page > task22_page else by_num[23]

        for num in range(1, 23):
            start_page, start_y = by_num[num]
            next_page, next_y = by_num[num + 1] if num < 22 else text2_start
            end_page, end_y = _task_end_before_next_marker(pdf, start_page, start_y, next_page, next_y)

            task_image = _cut_range(fitz_doc, pdf, start_page, start_y, end_page, end_y, end_margin=0)
            task_images[num] = _vstack([text1_image, task_image]) if num in (1, 2, 3) else task_image

        text2_start_page, text2_start_y = text2_start
        text2_image = _cut_range(fitz_doc, pdf, text2_start_page, text2_start_y, task23_page, task23_y, end_margin=0)

        for num in range(23, 28):
            start_page, start_y = by_num[num]
            if num < 27:
                next_page, next_y = by_num[num + 1]
                end_page, end_y = _task_end_before_next_marker(pdf, start_page, start_y, next_page, next_y)
            else:
                end_page = len(pdf.pages) - 1
                end_y = float(pdf.pages[end_page].height)
            task_image = _cut_range(fitz_doc, pdf, start_page, start_y, end_page, end_y, end_margin=0)
            task_images[num] = _vstack([text2_image, task_image]) if num in (23, 24, 25, 26) else task_image

        return task_images


def _text_to_png(text: str, width: int = 900) -> bytes:
    from PIL import ImageDraw, ImageFont

    padding = 30
    font_size = 17
    line_height = 24

    try:
        font = ImageFont.truetype("arial.ttf", font_size)
    except Exception:
        font = ImageFont.load_default()

    dummy = Image.new("RGB", (width, 100), "white")
    draw = ImageDraw.Draw(dummy)

    lines: List[str] = []
    current: List[str] = []
    for word in text.split():
        test = " ".join(current + [word])
        if current and draw.textlength(test, font=font) > width - padding * 2:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))

    height = max(200, len(lines) * line_height + padding * 2)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(lines):
        draw.text((padding, padding + index * line_height), line, fill="black", font=font)

    buf = io.BytesIO()
    image.save(buf, "PNG")
    return buf.getvalue()


def extract_task_images_docx(path: str) -> Dict[int, bytes]:
    if not DOCX_OK:
        raise ValueError("python-docx не установлен")

    doc = DocxDocument(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tasks_text: Dict[int, List[str]] = {}
    current: Optional[int] = None

    for paragraph in paragraphs:
        match = re.match(r"^(\d{1,2})[.\s]", paragraph)
        if match:
            num = int(match.group(1))
            if 1 <= num <= 27:
                current = num
                tasks_text[current] = [paragraph]
                continue
        if current is not None:
            tasks_text[current].append(paragraph)

    return {num: _text_to_png("\n".join(lines)) for num, lines in tasks_text.items()}


def extract_task_images(path: str) -> Dict[int, bytes]:
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return extract_task_images_pdf(path)
    if ext == "docx":
        return extract_task_images_docx(path)
    raise ValueError(f"Неподдерживаемый формат: {ext}")


def process_variant(variant_path: str, answers_path: str, tasks_dir: str) -> int:
    os.makedirs(tasks_dir, exist_ok=True)

    task_images = extract_task_images(variant_path)
    answers = parse_answers(answers_path)

    ordered_ids = list(range(1, 28))
    ordered_ids = [task_id for task_id in ordered_ids if task_id in task_images]

    saved_count = 0
    for task_id in ordered_ids:
        img_bytes = _compact_vertical_whitespace(task_images[task_id])
        img_bytes = _add_bottom_drawing_space(img_bytes)
        image_file = f"task_{task_id}.png"
        image_path = os.path.join(tasks_dir, image_file)
        with open(image_path, "wb") as file:
            file.write(img_bytes)

        task_type = TASK_TYPES.get(task_id, "unknown")
        if task_id == 0:
            display_id = "Текст 1"
        elif task_id == -1:
            display_id = "Текст 2"
        else:
            display_id = str(task_id)

        meta = {
            "id": task_id,
            "display_id": display_id,
            "type": task_type,
            "type_name": TASK_NAMES.get(task_type, "Задание"),
            "answer": answers.get(str(task_id), "") if task_id > 0 else "",
            "image_file": image_file,
            "is_text": task_id in (0, -1),
        }
        with open(os.path.join(tasks_dir, f"task_{task_id}_meta.json"), "w", encoding="utf-8") as file:
            json.dump(meta, file, ensure_ascii=False, indent=2)

        saved_count += 1

    correct = {
        str(num): {
            "answer": answers.get(str(num), ""),
            "type": TASK_TYPES.get(num, "unknown"),
        }
        for num in task_images
        if num > 0
    }
    with open(os.path.join(os.path.dirname(tasks_dir), "answers.json"), "w", encoding="utf-8") as file:
        json.dump(correct, file, ensure_ascii=False, indent=2)

    return saved_count
